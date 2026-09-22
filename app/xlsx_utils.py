"""
Some xlsx exporters (NREMT's roster export among them) emit a bare
`<fill/>` element in styles.xml with no `patternFill` child. Real Excel
tolerates this, but openpyxl's strict schema does not and raises
`TypeError: expected <class 'openpyxl.styles.fills.Fill'>` while loading
the stylesheet. Rather than ask users to hand-edit the file before every
upload, patch that one element in-memory and retry.
"""

import io
import re
import zipfile

import openpyxl

_BARE_FILL = re.compile(rb"<fill/>")
_BARE_FILL_REPLACEMENT = b'<fill><patternFill patternType="none"/></fill>'


def _repair_bare_fills(file_bytes):
    src = zipfile.ZipFile(io.BytesIO(file_bytes))
    out_buffer = io.BytesIO()
    with zipfile.ZipFile(out_buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/styles.xml":
                data = _BARE_FILL.sub(_BARE_FILL_REPLACEMENT, data)
            out.writestr(item, data)
    out_buffer.seek(0)
    return out_buffer


def load_workbook_safe(file_obj):
    """Load an xlsx from a file-like object, repairing known-bad but
    real-world stylesheet quirks if the first attempt fails."""
    file_bytes = file_obj.read()
    try:
        return openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except (TypeError, ValueError):
        pass

    try:
        repaired = _repair_bare_fills(file_bytes)
        return openpyxl.load_workbook(repaired, data_only=True)
    except Exception as exc:
        raise ValueError(
            "Could not read that spreadsheet, even after attempting to repair "
            "a known formatting quirk. Make sure it's a valid .xlsx file."
        ) from exc
